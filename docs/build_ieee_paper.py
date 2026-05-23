from pathlib import Path
import csv
import json
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
RESULTS = ROOT / "evaluation" / "results"
FIGURES = DOCS / "figures"
OUT = DOCS / "ShareRickshaw_IEEE_camera_ready_draft.docx"


TITLE = (
    "ShareRickshaw: Trust-Aware Grounded Route Planning for Informal "
    "Shared Autorickshaw Mobility in Mumbai"
)


def latest(pattern):
    files = sorted(RESULTS.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(pattern)
    return files[0]


def font(size=11, bold=False):
    try:
        return ImageFont.truetype("arial.ttf", size=size)
    except Exception:
        return ImageFont.load_default()


def rounded_box(draw, xy, text, fill, outline="#1F3A5F", text_color="#111111", width=2):
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    total_h = len(lines) * 18
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font(15, True))
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, y), line, fill=text_color, font=font(15, True))
        y += 18


def arrow(draw, start, end, label=None):
    draw.line([start, end], fill="#222222", width=3)
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 12
    p1 = (end[0] - size * math.cos(ang - math.pi / 6), end[1] - size * math.sin(ang - math.pi / 6))
    p2 = (end[0] - size * math.cos(ang + math.pi / 6), end[1] - size * math.sin(ang + math.pi / 6))
    draw.polygon([end, p1, p2], fill="#222222")
    if label:
        mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
        draw.text((mx - 42, my - 22), label, fill="#333333", font=font(12))


def draw_architecture(path):
    img = Image.new("RGB", (1400, 650), "white")
    d = ImageDraw.Draw(img)
    d.text((430, 28), "ShareRickshaw System Architecture", fill="#0B2545", font=font(30, True))
    rounded_box(d, (70, 210, 340, 370), "Passenger / Driver\nBrowser\nLeaflet + JS", "#E8F2FF")
    rounded_box(d, (560, 190, 850, 390), "Express / Node.js\nREST API + Socket.IO\nRoute Planner", "#EAF7EE")
    rounded_box(d, (1060, 110, 1320, 245), "MySQL\nusers, stands,\nroutes, bookings", "#FFF3D9")
    rounded_box(d, (1060, 315, 1320, 455), "Gemini API\nstructured route +\nplate extraction", "#F4EAFE")
    rounded_box(d, (560, 485, 850, 600), "Email / SOS\nEmergency Contacts", "#FDECEC")
    arrow(d, (340, 290), (560, 290), "REST / WS")
    arrow(d, (850, 220), (1060, 180), "SQL")
    arrow(d, (850, 340), (1060, 380), "JSON prompt")
    arrow(d, (700, 390), (700, 485), "SMTP")
    d.text((74, 525), "Frontend also uses OSRM for route visualisation.", fill="#444444", font=font(17))
    img.save(path)


def draw_pipeline(path):
    img = Image.new("RGB", (1200, 1200), "white")
    d = ImageDraw.Draw(img)
    d.text((332, 28), "Trust-Aware Grounded Ranking Pipeline", fill="#0B2545", font=font(28, True))
    boxes = [
        ("Origin, destination,\ndeparture time", "#E8F2FF"),
        ("Validate coordinates\nwithin Mumbai bounds", "#EAF7EE"),
        ("Enumerate MySQL\nstand-route pairs", "#EAF7EE"),
        ("Compute route features:\ntime, fare, access,\negress, evidence", "#FFF3D9"),
        ("Compute trust score\nT(c) and risk flags", "#FFF3D9"),
        ("Composite score S(c):\n0.42 time + 0.26 cost\n+ 0.18 access + 0.14 trust risk", "#F4EAFE"),
        ("Sort candidates\nascending by S(c)", "#EAF7EE"),
        ("Return ranked route\nwith evidence metadata", "#E8F2FF"),
    ]
    y = 100
    centers = []
    for text, fill in boxes:
        rounded_box(d, (320, y, 880, y + 92), text, fill)
        centers.append((600, y + 92))
        y += 130
    for i in range(len(centers) - 1):
        arrow(d, centers[i], (600, centers[i][1] + 38))
    rounded_box(d, (70, 462, 270, 565), "Gemini\noptional", "#FFFFFF", outline="#777777")
    rounded_box(d, (930, 462, 1130, 565), "Validate AI\nroute grounding", "#FFFFFF", outline="#777777")
    arrow(d, (320, 508), (270, 508), "prompt")
    arrow(d, (880, 508), (930, 508), "check")
    d.text((120, 1120), "LLM output is optional; deterministic ranking and fallback remain available.", fill="#444444", font=font(18))
    img.save(path)


def draw_scatter(csv_path, path):
    rows = []
    with open(csv_path, newline="") as f:
        for row in csv.DictReader(f):
            rows.append(
                {
                    "trust": float(row["trustScore"]),
                    "score": float(row["compositeScore"]),
                    "risk": row["hasRiskFlag"].lower() == "true",
                    "label": row["destination"],
                }
            )
    img = Image.new("RGB", (1200, 780), "white")
    d = ImageDraw.Draw(img)
    d.text((365, 30), "Trust Score vs Composite Route Score", fill="#0B2545", font=font(28, True))
    left, top, right, bottom = 120, 100, 1120, 650
    d.line((left, bottom, right, bottom), fill="#222222", width=3)
    d.line((left, top, left, bottom), fill="#222222", width=3)
    min_t, max_t = min(r["trust"] for r in rows), max(r["trust"] for r in rows)
    min_s, max_s = min(r["score"] for r in rows), max(r["score"] for r in rows)

    def x(v):
        return left + ((v - min_t) / (max_t - min_t)) * (right - left)

    def y(v):
        return bottom - ((v - min_s) / (max_s - min_s)) * (bottom - top)

    for tick in range(5):
        tx = left + tick * (right - left) / 4
        d.line((tx, bottom, tx, bottom + 8), fill="#222222", width=2)
        val = min_t + tick * (max_t - min_t) / 4
        d.text((tx - 18, bottom + 14), f"{val:.0f}", fill="#222222", font=font(14))
        sy = bottom - tick * (bottom - top) / 4
        d.line((left - 8, sy, left, sy), fill="#222222", width=2)
        sval = min_s + tick * (max_s - min_s) / 4
        d.text((left - 72, sy - 8), f"{sval:.2f}", fill="#222222", font=font(14))

    for r in rows:
        color = "#C43B3B" if r["risk"] else "#227A4B"
        cx, cy = x(r["trust"]), y(r["score"])
        d.ellipse((cx - 8, cy - 8, cx + 8, cy + 8), fill=color, outline="#222222")

    d.text((470, 705), "Trust score (higher is better)", fill="#222222", font=font(18, True))
    d.text((12, 350), "Composite", fill="#222222", font=font(16, True))
    d.text((18, 370), "score", fill="#222222", font=font(16, True))
    d.text((18, 390), "(lower better)", fill="#222222", font=font(16, True))
    d.ellipse((900, 120, 918, 138), fill="#227A4B", outline="#222222")
    d.text((930, 119), "No risk flags", fill="#222222", font=font(16))
    d.ellipse((900, 150, 918, 168), fill="#C43B3B", outline="#222222")
    d.text((930, 149), "Risk flag present", fill="#222222", font=font(16))
    img.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, True)
        set_cell_shading(hdr[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return table


def paragraph(doc, text="", style=None, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper() if level == 1 else text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10 if level == 1 else 9)
    r.bold = True
    return p


def add_math_problem_statement(doc):
    p = paragraph(doc)
    for text, italic, sub in [
        ("Given origin ", False, False),
        ("o", True, False),
        ("=(lat,lng), destination ", False, False),
        ("d", True, False),
        ("=(lat,lng), departure time ", False, False),
        ("t", True, False),
        (", and stand database ", False, False),
        ("D", True, False),
        ("={", False, False),
        ("s", True, False),
        ("1", False, True),
        (",...,", False, False),
        ("s", True, False),
        ("n", False, True),
        ("}, each stand ", False, False),
        ("s", True, False),
        ("i", False, True),
        (" contains fixed routes ", False, False),
        ("R", True, False),
        ("i", False, True),
        ("={", False, False),
        ("r", True, False),
        ("i1", False, True),
        (",...,", False, False),
        ("r", True, False),
        ("im", False, True),
        ("}. The output is an ordered list of route plans ", False, False),
        ("P", True, False),
        ("=[", False, False),
        ("p", True, False),
        ("1", False, True),
        (",", False, False),
        ("p", True, False),
        ("2", False, True),
        (",...]. The objective is to minimize ", False, False),
        ("S(c)", True, False),
        (" while exposing candidates where ", False, False),
        ("T(c)", True, False),
        (" < ", False, False),
        ("T", True, False),
        ("min", False, True),
        (".", False, False),
    ]:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
        r.italic = italic
        r.font.subscript = sub


def set_two_columns(section):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "360")


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    set_two_columns(section)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(9)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Author Name(s) Withheld for Final Submission\n"
        "Department of Computer Engineering, Bharatiya Vidya Bhavan's Sardar Patel Institute of Technology (S.P.I.T.), Mumbai, India"
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r.italic = True


def build():
    FIGURES.mkdir(exist_ok=True)
    scatter_csv = latest("trust_composite_scatter_*.csv")
    evidence_json = json.loads(latest("ieee_evidence_analysis_*.json").read_text())

    fig1 = FIGURES / "figure1_architecture.png"
    fig2 = FIGURES / "figure2_trust_ranking_pipeline.png"
    fig3 = FIGURES / "figure3_trust_composite_scatter.png"
    draw_architecture(fig1)
    draw_pipeline(fig2)
    draw_scatter(scatter_csv, fig3)

    doc = Document()
    style_doc(doc)
    add_title_block(doc)

    heading(doc, "Abstract", 1)
    paragraph(
        doc,
        "Shared autorickshaws provide low-cost last-mile mobility in Mumbai, but their fixed-corridor operations remain "
        "largely offline, creating uncertainty around route discovery, fare transparency, and passenger safety. This paper "
        "presents ShareRickshaw, a full-stack web prototype for informal paratransit routing, booking, driver coordination, "
        "and safety workflows. The central contribution is a trust-aware grounded route planner that ranks database-backed "
        "stand-route candidates using normalized travel time, fare, first/last-mile walking distance, and a transparent trust "
        "risk score. Optional LLM-generated routes are constrained to structured JSON and validated against known stands, "
        "destinations, transport modes, and cost-time totals before being returned. Offline seeded-data evaluation over seven "
        "Mumbai origin-destination cases achieved mean evidence recall of 1.0; the selected shared-auto options reduced "
        "estimated fare by 53% to 71% compared with direct-auto estimates, with 58% mean saving. Sensitivity analysis shows "
        "that top-ranked routes remain stable under most single-weight perturbations. The prototype demonstrates a reproducible "
        "approach for digitising informal shared mobility while making algorithmic evidence, limitations, and privacy requirements explicit.",
    )
    paragraph(doc, "Index Terms--informal transit, paratransit, route planning, shared autorickshaw, trust-aware ranking, LLM grounding.")

    heading(doc, "I. Introduction")
    paragraph(
        doc,
        "Shared autorickshaws operate as a low-cost, high-frequency, informal transit layer in many Indian cities. Unlike formal "
        "rail or bus services, their route knowledge is often local, corridor-based, and weakly represented in machine-readable "
        "feeds. A passenger may know a destination coordinate but not the correct stand, fixed route label, transfer point, or "
        "fare expectation. Generic routing tools can visualize roads, but they rarely prove that an informal shared-auto leg is "
        "connected to a real stand and corridor. ShareRickshaw addresses this gap through a grounded route-planning prototype "
        "for Mumbai-style shared autorickshaw travel.",
    )
    paragraph(
        doc,
        "The research challenge is not simply to build another booking interface. Informal shared-auto travel has three linked "
        "uncertainties: passengers may not know which physical stand serves their destination, fare information is often mediated "
        "through manual quotes, and first/last-mile walking exposure can change route acceptability, especially at night. A useful "
        "system therefore needs to expose evidence behind a recommendation instead of returning an opaque black-box route.",
    )
    heading(doc, "Formal Problem Statement", 2)
    add_math_problem_statement(doc)
    heading(doc, "Contributions", 2)
    paragraph(
        doc,
        "This paper makes four contributions: a full-stack prototype for shared-auto discovery, booking, driver coordination, fare "
        "estimation, and safety workflows; a grounded route-planning pipeline that validates optional LLM output against known stands "
        "and fixed routes; a trust-aware ranking function that exposes first/last-mile walking risk; and a reproducible evaluation "
        "package containing offline route cases, sensitivity analysis, fare-savings summaries, and figure generation scripts.",
    )

    heading(doc, "II. Related Work")
    paragraph(
        doc,
        "Prior work on informal transit digitisation, including the Digital Matatus project, shows that semi-formal transport "
        "systems can be made queryable through structured route data. GTFS offers a standard for formal transit feeds, but "
        "many shared-auto corridors lack complete schedules, stop sequences, and machine-readable fare rules. Recent work on "
        "retrieval-augmented generation and hallucination mitigation motivates a grounded approach for LLM use: the model may "
        "help synthesize route instructions, but database evidence and deterministic validation must constrain final output.",
    )
    paragraph(
        doc,
        "This work differs from conventional ride-hailing dispatch because it does not optimize assignment between a private driver "
        "and a single passenger. It reasons over fixed informal corridors represented by a stand-route registry. It also differs from "
        "unconstrained conversational route assistants because LLM output is optional, schema-bound, and rejected when it cannot be "
        "tied back to known database evidence.",
    )
    paragraph(
        doc,
        "The paper also connects to transport safety and data-governance literature. Mobility applications that collect location traces, "
        "emergency contacts, and vehicle identifiers can create value during distress events, but they also introduce retention, consent, "
        "and access-control responsibilities. For that reason, ShareRickshaw treats safety features as prototype workflows and explicitly "
        "separates them from claims of verified crime prevention or production-grade compliance.",
    )

    heading(doc, "III. System Architecture")
    paragraph(
        doc,
        "ShareRickshaw consists of a browser frontend, an Express/Node.js backend, MySQL-backed stand and route tables, Socket.IO "
        "booking notifications, and Gemini-assisted structured route and license-plate extraction. Fig. 1 summarizes the deployed "
        "prototype architecture.",
    )
    paragraph(
        doc,
        "The database separates user accounts, stands, fixed shared-auto routes, booking records, safety contacts, SOS logs, and "
        "night-tracking events. This separation supports reproducibility: the route planner can be evaluated offline from seed data "
        "without requiring live bookings, emergency events, or an active Gemini API key. The API returns route evidence metadata so "
        "each recommendation can be inspected during experiments.",
    )
    paragraph(
        doc,
        "The route-planning service is deliberately isolated from the controller layer. This makes the deterministic scoring and validation "
        "functions testable without an HTTP server, JWT token, MySQL connection, or Gemini key. That design choice is useful for research "
        "reproducibility because offline evaluation scripts can import the same ranking functions used by the live API.",
    )
    doc.add_picture(str(fig1), width=Inches(3.15))
    caption(doc, "Fig. 1. ShareRickshaw system architecture.")

    heading(doc, "IV. Trust-Aware Grounded Route Planning")
    paragraph(
        doc,
        "The planner first validates coordinates, loads known stand-route pairs, and computes deterministic candidates. Each "
        "candidate is represented by travel time, fare, access walking distance, egress walking distance, and evidence completeness. "
        "The trust score penalizes long first/last-mile walking, with extra exposure under night-travel conditions. The composite "
        "ranking score is:",
    )
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("S(c) = 0.42 time + 0.26 cost + 0.18 access + 0.14 trustRisk")
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.italic = True
    paragraph(
        doc,
        "Lower values of S(c) are preferred. The current weights are design weights stress-tested by sensitivity analysis rather "
        "than claimed as population-optimal behavioral coefficients. Fig. 2 shows the route-ranking pipeline and LLM validation branch.",
    )
    doc.add_picture(str(fig2), width=Inches(3.15))
    caption(doc, "Fig. 2. Trust-aware grounded ranking and validation pipeline.")

    heading(doc, "Trust Score", 2)
    paragraph(
        doc,
        "The trust score is intentionally transparent. It begins from an evidence-backed default, adds confidence when the stand, "
        "fixed route, and coordinates are present, and subtracts penalties for long access or egress walking. Under night-travel "
        "conditions, candidates with long first/last-mile exposure receive an additional penalty. The score should not be interpreted "
        "as a crime prediction or formal safety guarantee; it is a user-facing indicator of route evidence quality and walking exposure.",
    )
    paragraph(
        doc,
        "This design avoids two common failure modes. First, it prevents low-cost but uncomfortable routes from being silently promoted when "
        "they involve long walking exposure after the shared-auto leg. Second, it prevents LLM-generated instructions from appearing authoritative "
        "when they mention unsupported stands, unsupported transport modes, or inconsistent totals. The planner can still return a direct-auto "
        "estimate and deterministic fallback, so service degradation is explicit rather than hidden.",
    )

    heading(doc, "Algorithm", 2)
    add_table(
        doc,
        ["Step", "Operation"],
        [
            ["1", "Validate origin and destination within supported Mumbai prototype bounds."],
            ["2", "Load stand-route registry from MySQL or seeded offline data."],
            ["3", "Generate deterministic direct-auto estimate for comparison."],
            ["4", "Enumerate stand-route candidates and compute time, fare, access, and egress features."],
            ["5", "Compute trust score and composite ranking score for each candidate."],
            ["6", "Optionally request a structured Gemini route using the stand-route registry."],
            ["7", "Validate generated route steps against known stands, destinations, modes, and totals."],
            ["8", "Return ranked deterministic, AI-validated, or fallback route options with evidence metadata."],
        ],
        [0.35, 2.65],
    )
    caption(doc, "Table I. Grounded structured multimodal planning procedure.")

    heading(doc, "V. Implementation")
    paragraph(
        doc,
        "The backend implements deterministic route estimation, candidate scoring, trust-aware re-ranking, structured Gemini prompting, "
        "and post-generation validation. The validator checks allowed transport modes, known stand names, known route destinations, "
        "and consistency between step-level and route-level cost/time totals. If the AI route is unavailable or weakly grounded, the "
        "system returns deterministic fallback routes with evidence metadata. The frontend exposes route options, fare estimation, "
        "booking flows, SOS alerts, emergency contacts, and night-tracking workflows.",
    )
    paragraph(
        doc,
        "The API response includes evidence metadata such as planner type, stand identifier, route identifier, candidate score, walking distances, "
        "trust score, and validation warnings. These fields are not merely debugging details; they are the mechanism by which the paper's claims "
        "can be audited. A reviewer can trace a displayed recommendation back to a structured database entry and to the exact ranking logic that "
        "selected it.",
    )
    add_table(
        doc,
        ["Module", "Implementation role"],
        [
            ["Authentication", "JWT-based login flows with separate user and driver-facing screens."],
            ["Stand registry", "MySQL tables and seed data for stands, destinations, fares, times, and coordinates."],
            ["Route planner", "Deterministic scoring, trust-aware ranking, Gemini prompting, validation, and fallback routing."],
            ["Booking", "Passenger request flow and driver dashboard supported by REST and Socket.IO."],
            ["Safety", "Emergency contacts, SOS email flow, night-tracking records, and ALPR-assisted plate extraction."],
            ["Evaluation", "Offline route-case runner, regression tests, sensitivity sweep, and generated paper artifacts."],
        ],
        [0.85, 2.15],
    )
    caption(doc, "Table II. Prototype modules and implementation roles.")

    heading(doc, "VI. Evaluation")
    paragraph(
        doc,
        "Evaluation uses offline seeded-data experiments because live MySQL credentials and production deployment were not available "
        "during this study. The seed database contains 10 stands and 31 fixed route entries. Seven origin-destination cases were chosen "
        "to cover distinct Mumbai localities and risk combinations.",
    )
    paragraph(
        doc,
        "The evaluation is therefore a controlled route-grounding study rather than a city-wide deployment. The main dependent variables "
        "are evidence recall, fare savings relative to a deterministic direct-auto baseline, top-ranked candidate changes under trust-aware "
        "ranking, and sensitivity of the ranking to weight perturbation. This framing avoids overclaiming while still testing whether the "
        "implemented algorithm behaves reproducibly.",
    )
    paragraph(
        doc,
        "Evidence recall is computed by checking whether the selected and fallback route descriptions contain the expected stand and destination "
        "labels for each seeded case. Fare savings are calculated by comparing the selected grounded shared-auto candidate against the deterministic "
        "direct-auto estimate produced by the same service. The sensitivity sweep varies one ranking weight at a time from 0.0 to 0.5 while rescaling "
        "the remaining weights proportionally so the total remains 1.0.",
    )
    add_table(
        doc,
        ["Property", "Coverage"],
        [
            ["Zero access walk candidate", "7/7"],
            ["Non-zero egress in selected shared-auto route", "3/7"],
            ["Night travel", "2/7"],
            ["Long egress candidate flagged (>1.2 km)", "6/7"],
            ["Ranking changed by trust-aware composite", "3/7"],
            ["Distinct Mumbai localities represented", "7/7"],
        ],
        [2.15, 0.8],
    )
    caption(doc, "Table III. Seven-case coverage summary.")
    paragraph(
        doc,
        "The planner achieved mean evidence recall of 1.0 against expected stand and destination labels. The selected shared-auto "
        f"options produced fare savings from {round(evidence_json['fareSavings']['min']*100)}% to {round(evidence_json['fareSavings']['max']*100)}%, "
        f"with {round(evidence_json['fareSavings']['mean']*100)}% mean saving relative to direct-auto estimates.",
    )
    add_table(
        doc,
        ["Case", "Direct", "Shared", "Saving"],
        [
            ["Bandra", "Rs 34", "Rs 15", "56%"],
            ["Andheri", "Rs 51", "Rs 15", "71%"],
            ["Dadar", "Rs 36", "Rs 15", "58%"],
            ["Kurla", "Rs 23", "Rs 10", "57%"],
            ["Malad", "Rs 23", "Rs 10", "57%"],
            ["Powai night", "Rs 43", "Rs 20", "53%"],
            ["Ghatkopar night", "Rs 57", "Rs 25", "56%"],
        ],
        [0.95, 0.55, 0.55, 0.55],
    )
    caption(doc, "Table IV. Fare savings for selected shared-auto options.")
    doc.add_picture(str(fig3), width=Inches(3.15))
    caption(doc, "Fig. 3. Trust score versus composite route score for the top three candidates in each evaluation case.")

    heading(doc, "Weight Sensitivity", 2)
    add_table(
        doc,
        ["Weight", "Stable range", "Observed top-1 changes"],
        [
            ["time", "0.0-0.5", "0/7 at all sweep points"],
            ["cost", "0.2-0.4", "changes at 0.0, 0.1, and 0.5"],
            ["access", "0.0-0.3", "changes at 0.4 and 0.5"],
            ["trust", "0.0-0.5", "0/7 at all sweep points"],
        ],
        [0.55, 0.75, 1.65],
    )
    caption(doc, "Table V. Summary of one-at-a-time sensitivity sweep.")
    heading(doc, "Bandra Case Interpretation", 2)
    paragraph(
        doc,
        "Both Bandra candidates have identical trust scores (96.0), so the ranking change is driven by cost-time trade-off rather "
        "than trust penalisation. Trust-aware ranking selected Linking Road because its fare advantage (Rs 15 versus Rs 25) dominated "
        "its small time disadvantage (16 versus 14 min) after normalization.",
    )
    paragraph(
        doc,
        "This is an important qualitative result because it shows that the composite function is not simply a safety filter. In cases where trust "
        "is equal, the ranking behaves as a multi-objective optimiser over cost, time, and walking distance. In cases with long egress or night "
        "walking, the trust term makes the risk visible and can prevent superficially cheap options from appearing unqualified.",
    )
    heading(doc, "Preliminary Usability Evaluation Protocol", 2)
    paragraph(
        doc,
        "A human-factor evaluation instrument is prepared for final submission. It asks 5-10 participants to rate route understandability, "
        "fare transparency, and perceived safety from SOS/night tracking on a five-point Likert scale. Participant statistics are not "
        "reported here because responses must be collected from real users before submission. The repository includes the questionnaire "
        "so that these results can be added without changing the technical evaluation pipeline.",
    )
    paragraph(
        doc,
        "Once responses are collected, the intended report format is mean plus standard deviation for each item. The three items measure whether "
        "participants understand the recommended route and transfer points, whether fare transparency improves trust compared with manual quotes, "
        "and whether SOS and night tracking increase perceived night-travel safety. This keeps the human-factor study lightweight but aligned with "
        "the system's actual claims.",
    )

    heading(doc, "VII. Privacy and Data Governance")
    paragraph(
        doc,
        "The prototype processes sensitive mobility and safety data, including GPS coordinates, booking timestamps, emergency contacts, "
        "and license-plate text extracted from user-submitted images. Night tracking is user-initiated and should be terminable by the "
        "passenger. Production deployment would require explicit consent flows, data minimisation, a retention schedule, access controls, "
        "and review against India's Digital Personal Data Protection Act, 2023.",
    )

    heading(doc, "VIII. Limitations")
    paragraph(
        doc,
        "The current evaluation is offline and seed-data based; it does not prove city-wide coverage, real-time traffic optimality, official "
        "tariff compliance, production reliability, or passenger safety impact. ALPR accuracy and user-perceived usability require labeled "
        "image data and real participant responses before final submission.",
    )
    paragraph(
        doc,
        "The seven cases are deliberately diverse but not statistically exhaustive. The ranking weights are stress-tested but not derived "
        "from a large stated-preference survey. These constraints make the paper strongest as a controlled prototype and algorithmic "
        "evaluation, with live deployment, larger route datasets, and human-subject validation left for the next study phase.",
    )
    paragraph(
        doc,
        "A production system would also need live feed maintenance, route closure handling, official fare integration, accessibility review, "
        "driver-side misuse controls, and a formal privacy impact assessment. These requirements are outside the current controlled prototype, "
        "but the architecture leaves room for them through explicit route evidence, modular safety tables, and reproducible evaluation scripts.",
    )

    heading(doc, "IX. Conclusion")
    paragraph(
        doc,
        "ShareRickshaw demonstrates a grounded, trust-aware route-planning prototype for informal shared-auto mobility. Its main contribution "
        "is not LLM usage alone, but the evidence layer around route generation: structured stand-route data, deterministic ranking, transparent "
        "trust features, validation of AI output, and reproducible offline evaluation. This makes the system a credible base for controlled "
        "pilot studies and future larger-scale informal transit digitisation.",
    )

    heading(doc, "References")
    refs = [
        "[1] S. Williams, A. White, P. Waiganjo, D. Orwa, and J. M. Klopp, \"The digital matatu project: Using cell phones to create open source data for Nairobi's semi-formal bus system,\" Journal of Transport Geography, 2015.",
        "[2] General Transit Feed Specification, \"GTFS documentation overview,\" MobilityData.",
        "[3] Y. Gao et al., \"Retrieval-Augmented Generation for Large Language Models: A Survey,\" arXiv:2312.10997.",
        "[4] L. Huang et al., \"A Survey on Hallucination in Large Language Models: Principles, Taxonomy, Challenges, and Open Questions,\" ACM Transactions on Information Systems, 2025.",
        "[5] R. Laroca et al., \"A Robust Real-Time Automatic License Plate Recognition Based on the YOLO Detector,\" arXiv:1802.09567.",
        "[6] Government of India, \"The Digital Personal Data Protection Act, 2023,\" India Code.",
    ]
    for ref in refs:
        paragraph(doc, ref)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
